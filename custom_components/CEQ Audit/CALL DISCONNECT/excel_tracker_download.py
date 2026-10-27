import io
import re
import json
import base64
import asyncio
import urllib.parse
from pathlib import Path
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import requests

from lfx.custom.custom_component.component import Component
from lfx.io import (
    DataInput,
    StrInput,
    SecretStrInput,
    DropdownInput,
    IntInput,
    BoolInput,
    Output,
)
from lfx.schema.data import Data
from lfx.schema.message import Message

ET = ZoneInfo("America/New_York")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
STAMP_RE = re.compile(r"(\d{8}_\d{6})\.xlsx$")
ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
BAD_TITLE = re.compile(r"[\[\]\:\*\?\/\\]")

SHEET_ORDER = ["Tracker", "Evidence Timeline", "Agent Summary", "Settings"]
VERDICT_ORDER = ["LIKELY_WORK_AVOIDANCE", "POSSIBLE_WORK_AVOIDANCE", "NOT_WORK_AVOIDANCE",
                 "TECHNICAL_OR_DATA_ISSUE", "INSUFFICIENT_DATA", "NOT_REVIEWED", "REVIEW_FAILED"]
VERDICT_RANK = {v: i for i, v in enumerate(VERDICT_ORDER)}
VERDICT_FILL = {
    "LIKELY_WORK_AVOIDANCE": "FFC7CE", "POSSIBLE_WORK_AVOIDANCE": "FFEB9C",
    "NOT_WORK_AVOIDANCE": "C6EFCE", "TECHNICAL_OR_DATA_ISSUE": "DDEBF7",
    "INSUFFICIENT_DATA": "EDEDED", "NOT_REVIEWED": "EDEDED", "REVIEW_FAILED": "F4B084",
}
RED, AMBER, HEADER = "FFC7CE", "FFEB9C", "1F4E78"
WRAP_COLS = {"AI Review", "Behaviours Observed", "Legitimate Explanations", "Coaching Points", "QA Check", "Review Error", "Unrelated SNOW Detail", "Rule Flags", "Event", "Reviewer Comments", "Value", "Meaning", "Audit Reasons", "SNOW Tickets Touched", "Transcript (opening)", "Email vs Call", "Channel Note", "Email Start (ET)", "Email End (ET)", "Email Interaction(s)"}
YES_COLS = {"Call Left Unspoken", "Email Actioned During Call"}
DROPDOWNS = {
    "Reviewer Decision": '"Confirmed Work Avoidance,Not Work Avoidance,Technical Issue,Needs More Info"',
    "Status": '"Open,In Review,Closed"',
}
FREEZE = {"Tracker": "D2", "Evidence Timeline": "D2"}
ATTENTION_COLS = ["Answered (ET)", "Agent", "Conversation ID", "AI Verdict", "AI Risk Score",
                  "Recommended Action", "QA Check"]

LEGEND = [
    ("Verdict", "LIKELY_WORK_AVOIDANCE", "Clear evidence the agent was connected to an engaging caller and did not engage, "
     "or did unrelated SNOW/email work while the caller was left unattended, with no legitimate explanation."),
    ("Verdict", "POSSIBLE_WORK_AVOIDANCE", "Some indicators present but evidence incomplete or mixed. Needs human review."),
    ("Verdict", "NOT_WORK_AVOIDANCE", "Indicators fully explained (own ticket, grace period, hold, internal call)."),
    ("Verdict", "TECHNICAL_OR_DATA_ISSUE", "Evidence points to audio, recording, channel-mapping or data problems."),
    ("Verdict", "INSUFFICIENT_DATA", "Key data missing; no conclusion possible."),
    ("Verdict", "NOT_REVIEWED", "Not sent to the AI (review cap, time budget or Matcha limit). See Review Error."),
    ("Verdict", "REVIEW_FAILED", "AI call failed or returned invalid output. See Review Error."),
    ("Flag", "SILENT_ANSWER", "Agent channel had (almost) no speech while connected."),
    ("Flag", "BOTH_SILENT", "Neither side spoke: ghost call, one-way audio or recording fault."),
    ("Flag", "DELAYED_GREETING", "First agent speech came well after the answer."),
    ("Flag", "CALLER_UNANSWERED", "Caller tried to engage (hello? / are you there?) before the agent spoke."),
    ("Flag", "AGENT_RELEASED_SILENT_CALL", "Agent side disconnected a call in which the agent did not speak."),
    ("Flag", "LONG_AGENT_SILENCE", "Long agent silence outside hold during an otherwise engaged call."),
    ("Flag", "SNOW_ACTIVITY_DURING_CALL", "Unrelated, non-grace ServiceNow email/work note/update while connected."),
    ("Flag", "GENESYS_EMAIL_DURING_CALL", "A Genesys email interaction was worked during the call."),
    ("Flag", "MONO_AUDIO", "Single mixed channel: speech cannot be attributed to agent or caller."),
    ("Flag", "NO_AUDIO_EVIDENCE", "Recording missing, archived, deleted or unreadable."),
    ("Column", "Call Left Unspoken", "Computed from audio: YES = silent answer, PARTIAL = delayed greeting or dead air, "
     "NO = engaged normally, UNKNOWN = no reliable audio."),
    ("Column", "Email Actioned During Call", "Computed from SNOW/Genesys data: YES = unrelated email or SNOW work "
     "during the call, NO = none, UNKNOWN = SNOW data not loaded."),
    ("Column", "QA Check", "Consistency check of the AI answer against the measured data. Review rows that are not OK first."),
    ("Column", "Pre-Score / AI Risk Score", "Pre-Score is a rule heuristic. AI Risk Score: 0-20 none, 21-50 low, "
     "51-75 moderate, 76-100 high."),
    ("Column", "Audio Evidence", "COMPLETE = voice activity + opening transcript; VAD_ONLY = voice activity only; "
     "other values = no usable audio."),
    ("Sheet", "Evidence Timeline", "Origin = Code rows are measured facts; Origin = AI rows are the mission's narrative."),
    ("Note", "Advisory", "AI output is advisory. A human reviewer makes the final decision."),
    ("Flag", "EMAIL_ACTIONED_DURING_CALL", "A Genesys email was opened or finished while the call was connected."),
    ("Flag", "SYSTEM_PROMPT_ONLY", "The only speech on the agent side was IVR/system audio; the agent did not speak."),
    ("Flag", "DISCONNECT_ERROR", "Genesys recorded the disconnect as error/system. A platform fault cannot be ruled out."),
    ("Flag", "CHECK_CHANNEL_MAPPING", "Agent/caller tracks may be swapped; check the Transcript column."),
    ("Column", "Transcript (opening)", "Lines labelled AGENT, CALLER or SYSTEM (IVR) with the recording track number."),
    ("Column", "Email vs Call", "How each Genesys email interaction's start and end relate to the call."),
]


# ---------------- helpers ----------------
def _data(v):
    if isinstance(v, list) and v:
        v = v[0]
    if hasattr(v, "data"):
        v = v.data or {}
    return v if isinstance(v, dict) else {}


def _safe(s, limit=40):
    s = re.sub(r"[^A-Za-z0-9]+", "_", str(s or "")).strip("_")
    return s[:limit]


def _int(value, default):
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _num(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _expired(name, keep_days):
    m = STAMP_RE.search(name)
    if not m or keep_days <= 0:
        return False
    try:
        made = datetime.strptime(m.group(1), "%Y%m%d_%H%M%S")
    except ValueError:
        return False
    return datetime.now() - made > timedelta(days=keep_days)


def _title(name):
    return BAD_TITLE.sub("_", str(name or "Sheet"))[:31] or "Sheet"


def _put(ws, row, col, value):
    """Write one cell safely (no formulas, no illegal characters, Excel length limit)."""
    if value is None or (isinstance(value, float) and value != value):
        value = ""
    elif isinstance(value, (list, tuple, set)):
        value = "; ".join(str(x) for x in value if x not in (None, ""))
    elif isinstance(value, dict):
        value = json.dumps(value, ensure_ascii=False)
    elif not isinstance(value, (str, int, float, bool)):
        value = str(value)
    if isinstance(value, str):
        value = ILLEGAL_XML.sub("", value)[:32000]
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # keep as text, never a formula
    return cell


def _sheets_from(payload):
    notes, out = [], []
    sheets = payload.get("sheets")
    if isinstance(sheets, dict) and sheets:
        names = [n for n in SHEET_ORDER if n in sheets] + [n for n in sheets if n not in SHEET_ORDER]
        for n in names:
            s = sheets.get(n) or {}
            rows = [r for r in s.get("rows") or [] if isinstance(r, dict)]
            cols = [str(c) for c in s.get("columns") or []] or (list(rows[0].keys()) if rows else [])
            out.append((n, cols, rows))
    else:
        rows = [r for r in payload.get("rows") or [] if isinstance(r, dict)]
        out.append(("Tracker", list(rows[0].keys()) if rows else [], rows))
        notes.append("Reviewer & Tracker not patched: only the Tracker sheet is available "
                     "(apply Patch B and C to get Timeline, Summary and Settings).")
    return out, notes


def _verdict_counts(payload, tracker_rows):
    v = payload.get("verdicts")
    if isinstance(v, dict) and v:
        return {str(k): _int(n, 0) for k, n in v.items()}
    out = {}
    for r in tracker_rows:
        k = str(r.get("AI Verdict") or "")
        if k:
            out[k] = out.get(k, 0) + 1
    return out


def _qa_issue(r):
    q = str(r.get("QA Check") or "")
    return bool(q) and q != "OK" and not q.startswith("Auto")


def _attention(rows, limit):
    pick = [r for r in rows if r.get("AI Verdict") in ("LIKELY_WORK_AVOIDANCE", "POSSIBLE_WORK_AVOIDANCE")
            or _qa_issue(r)]
    pick.sort(key=lambda r: (VERDICT_RANK.get(str(r.get("AI Verdict")), 9),
                             -_num(r.get("AI Risk Score")), str(r.get("Answered (ET)") or "")))
    return pick[:limit]


def _agents(payload, tracker_rows):
    a = payload.get("agents")
    if isinstance(a, list) and a:
        return [str(x) for x in a if x]
    return sorted({str(r.get("Agent")) for r in tracker_rows if r.get("Agent")})


# ---------------- workbook ----------------
def _data_sheet(wb, name, cols, rows, url_template):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    ws = wb.create_sheet(_title(name))
    if not cols:
        ws.cell(row=1, column=1, value="No data for this sheet in this run.")
        return 0

    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=HEADER)
    wrap, top = Alignment(wrap_text=True, vertical="top"), Alignment(vertical="top")
    link_font, ai_font = Font(color="0563C1", underline="single"), Font(italic=True, color="595959")
    fills = {k: PatternFill("solid", fgColor=v) for k, v in VERDICT_FILL.items()}
    red, amber = PatternFill("solid", fgColor=RED), PatternFill("solid", fgColor=AMBER)

    for ci, col in enumerate(cols, 1):
        c = ws.cell(row=1, column=ci, value=col)
        c.font, c.fill, c.alignment = head_font, head_fill, wrap
    idx = {c: i + 1 for i, c in enumerate(cols)}

    for ri, row in enumerate(rows, 2):
        for ci, col in enumerate(cols, 1):
            cell = _put(ws, ri, ci, row.get(col, ""))
            cell.alignment = wrap if col in WRAP_COLS else top
            if name == "Evidence Timeline" and row.get("Origin") == "AI":
                cell.font = ai_font
        verdict = str(row.get("AI Verdict") or "")
        if "AI Verdict" in idx and verdict in fills:
            ws.cell(row=ri, column=idx["AI Verdict"]).fill = fills[verdict]
        for col in YES_COLS:
            if col in idx:
                v = str(row.get(col) or "").upper()
                if v == "YES":
                    ws.cell(row=ri, column=idx[col]).fill = red
                elif v == "PARTIAL":
                    ws.cell(row=ri, column=idx[col]).fill = amber
        if "QA Check" in idx and _qa_issue(row):
            ws.cell(row=ri, column=idx["QA Check"]).fill = amber
        if "Speaker Mapping Warning" in idx and str(row.get("Speaker Mapping Warning")).lower() == "true":
            ws.cell(row=ri, column=idx["Speaker Mapping Warning"]).fill = amber
        if name == "Agent Summary":
            for col, fill in (("LIKELY_WORK_AVOIDANCE", red), ("POSSIBLE_WORK_AVOIDANCE", amber)):
                if col in idx and _int(row.get(col), 0) > 0:
                    ws.cell(row=ri, column=idx[col]).fill = fill
        if url_template and "Conversation ID" in idx and row.get("Conversation ID"):
            cell = ws.cell(row=ri, column=idx["Conversation ID"])
            cell.hyperlink = url_template.replace("{id}", str(row.get("Conversation ID")))
            cell.font = link_font

    ws.freeze_panes = FREEZE.get(name, "A2")
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for ci, col in enumerate(cols, 1):
        letter = get_column_letter(ci)
        if col in WRAP_COLS:
            ws.column_dimensions[letter].width = 60
        else:
            longest = max([len(str(col))] + [len(str(r.get(col, ""))) for r in rows[:300]])
            ws.column_dimensions[letter].width = min(40, max(10, longest + 2))
    for col, formula in DROPDOWNS.items():
        if col in idx:
            letter = get_column_letter(idx[col])
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(letter + "2:" + letter + str(max(500, len(rows) + 200)))
    return len(rows)


def _overview(wb, payload, tracker_rows, notes, attention_limit):
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.create_sheet("Overview", 0)
    bold = Font(bold=True)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=HEADER)
    wrap = Alignment(wrap_text=True, vertical="top")

    ws["A1"] = "Work Avoidance Tracker"
    ws["A1"].font = Font(bold=True, size=16, color=HEADER)
    ws["A2"] = "Advisory AI review of completed, recorded calls. A human reviewer makes the final decision."
    ws["A2"].font = Font(italic=True, color="595959")

    agents = _agents(payload, tracker_rows)
    counts = _verdict_counts(payload, tracker_rows)
    facts = [
        ("Window", payload.get("interval", "")),
        ("Agents", ", ".join(agents) or "none"),
        ("Answered calls analysed", payload.get("answered_calls", "")),
        ("Flagged calls", payload.get("flagged_calls", len(tracker_rows))),
        ("AI-reviewed calls", payload.get("ai_reviewed", sum(counts.get(v, 0) for v in VERDICT_ORDER[:5]))),
        ("QA checks raised", payload.get("qa_issues", sum(1 for r in tracker_rows if _qa_issue(r)))),
        ("Matcha copy", (str(payload.get("file_name") or "") + (" (file id " + str(payload.get("file_id")) + ")"
                         if payload.get("file_id") else "")) if payload.get("uploaded_to_matcha") else "not uploaded"),
        ("Generated (ET)", datetime.now(ET).strftime("%m/%d/%Y %H:%M:%S")),
    ]
    r = 4
    for k, v in facts:
        ws.cell(row=r, column=1, value=k).font = bold
        _put(ws, r, 2, v).alignment = wrap
        r += 1

    r += 1
    for ci, h in enumerate(["Verdict", "Calls"], 1):
        c = ws.cell(row=r, column=ci, value=h)
        c.font, c.fill = head_font, head_fill
    r += 1
    for v in VERDICT_ORDER + [k for k in counts if k not in VERDICT_ORDER]:
        ws.cell(row=r, column=1, value=v).fill = PatternFill("solid", fgColor=VERDICT_FILL.get(v, "FFFFFF"))
        ws.cell(row=r, column=2, value=counts.get(v, 0))
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Review these first (LIKELY / POSSIBLE verdicts and QA issues)").font = Font(
        bold=True, size=12, color=HEADER)
    r += 1
    for ci, h in enumerate(ATTENTION_COLS, 1):
        c = ws.cell(row=r, column=ci, value=h)
        c.font, c.fill = head_font, head_fill
    r += 1
    top = _attention(tracker_rows, attention_limit)
    if not top:
        ws.cell(row=r, column=1, value="No LIKELY or POSSIBLE verdicts and no QA issues in this run.")
        r += 1
    for row in top:
        for ci, h in enumerate(ATTENTION_COLS, 1):
            cell = _put(ws, r, ci, row.get(h, ""))
            cell.alignment = wrap
            if h == "AI Verdict" and str(row.get(h)) in VERDICT_FILL:
                cell.fill = PatternFill("solid", fgColor=VERDICT_FILL[str(row.get(h))])
        r += 1

    warnings = [str(w) for w in payload.get("warnings") or [] if w] + list(notes)
    if warnings:
        r += 1
        ws.cell(row=r, column=1, value="Warnings").font = Font(bold=True, color="C00000")
        r += 1
        for w in warnings:
            _put(ws, r, 1, w).alignment = wrap
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
            r += 1

    r += 1
    ws.cell(row=r, column=1, value="How to review").font = Font(bold=True, color=HEADER)
    steps = [
        "1. Tracker sheet: filter QA Check not equal to OK first - the AI disagreed with the measured data.",
        "2. Then filter AI Verdict = LIKELY_WORK_AVOIDANCE, then POSSIBLE_WORK_AVOIDANCE.",
        "3. Use Evidence Timeline (filter by Conversation ID) to see exact times of speech, holds and SNOW actions.",
        "4. Listen to the recording for any call you intend to act on, then fill Reviewer, "
        "Reviewer Decision and Status.",
        "5. TECHNICAL_OR_DATA_ISSUE with CHECK_TELEPHONY should go to the telephony team, not performance review.",
    ]
    for s in steps:
        r += 1
        _put(ws, r, 1, s)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)

    for letter, width in zip("ABCDEFG", (34, 40, 40, 28, 14, 22, 60)):
        ws.column_dimensions[letter].width = width


def _legend(wb):
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.create_sheet("Legend")
    for ci, h in enumerate(["Type", "Name", "Meaning"], 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=HEADER)
    for ri, (t, n, m) in enumerate(LEGEND, 2):
        ws.cell(row=ri, column=1, value=t)
        cell = ws.cell(row=ri, column=2, value=n)
        if n in VERDICT_FILL:
            cell.fill = PatternFill("solid", fgColor=VERDICT_FILL[n])
        ws.cell(row=ri, column=3, value=m).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 100
    ws.freeze_panes = "A2"


def _build_workbook(payload, sheets, notes, url_template, overview, legend, attention_limit):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    counts = {}
    for name, cols, rows in sheets:
        counts[name] = _data_sheet(wb, name, cols, rows, url_template)
    if overview:
        tracker_rows = next((rows for n, c, rows in sheets if n == "Tracker"), [])
        _overview(wb, payload, tracker_rows, notes, attention_limit)
    if legend:
        _legend(wb)
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), counts


# ---------------- component ----------------
class CustomComponent(Component):
    display_name = "Excel Tracker Download"
    description = ("Builds the formatted Work Avoidance Excel tracker from the Reviewer's Tracker Data, saves it in "
                   "the flow file store and returns a download link.")
    icon = "file-spreadsheet"
    name = "ExcelTrackerDownload"

    inputs = [
        DataInput(name="tracker_data", display_name="Tracker Data", required=True,
                  info="Connect Work Avoidance Reviewer & Tracker > Tracker Data."),
        StrInput(name="file_prefix", display_name="File Prefix", value="Work_Avoidance_Tracker"),
        DropdownInput(name="storage", display_name="Storage",
                      options=["flow file store (download link)", "flow file store + local copy", "local folder only"],
                      value="flow file store (download link)"),
        StrInput(name="base_url", display_name="MDS Base URL", value="https://matchaflow.harriscomputer.com"),
        SecretStrInput(name="flow_api_key", display_name="MDS API Key (fallback upload)", required=False,
                       info="Only used if the internal storage service is unavailable."),
        IntInput(name="keep_days", display_name="Delete Trackers Older Than (days)", value=7,
                 info="Trackers contain agent performance and caller data. 0 = never delete."),
        BoolInput(name="include_overview", display_name="Include Overview Sheet", value=True),
        BoolInput(name="include_legend", display_name="Include Legend Sheet", value=True),
        IntInput(name="attention_rows", display_name="Overview 'Review First' Rows", value=25, advanced=True),
        StrInput(name="interaction_url_template", display_name="Conversation Link Template", value="",
                 advanced=True,
                 info="Optional. Makes Conversation ID clickable. Use {id} as placeholder, e.g. your Genesys "
                      "interaction details URL copied from the browser with the ID replaced by {id}."),
        StrInput(name="local_dir", display_name="Local Folder", value="/app/generated_reports", advanced=True),
        BoolInput(name="inline_link", display_name="Add Inline Data Link", value=False, advanced=True,
                  info="Embeds the file as a data: link. Some chat views block these."),
        IntInput(name="inline_max_kb", display_name="Inline Link Max Size (KB)", value=500, advanced=True),
    ]

    outputs = [
        Output(display_name="Download Link", name="download", method="build_download", group_outputs=True),
        Output(display_name="File Info", name="file_info", method="build_info", group_outputs=True),
    ]

    # ---------------- storage ----------------
    def _flow_id(self):
        fid = getattr(self, "flow_id", None)
        if not fid:
            fid = getattr(getattr(self, "graph", None), "flow_id", None)
        return str(fid or "")

    def _storage_service(self, errors):
        for path in ("lfx.services.deps", "langflow.services.deps"):
            try:
                mod = __import__(path, fromlist=["get_storage_service"])
                svc = mod.get_storage_service()
                if svc is not None:
                    return svc
            except Exception as e:
                errors.append(path + ": " + str(e)[:100])
        return None

    async def _save_store(self, flow_id, name, data, errors):
        svc = self._storage_service(errors)
        if svc is not None:
            try:
                await svc.save_file(flow_id=flow_id, file_name=name, data=data)
                await self._cleanup_store(svc, flow_id, errors)
                return "storage service", name
            except Exception as e:
                errors.append("storage service save: " + str(e)[:150])
        key = str(self.flow_api_key or "").strip()
        if not key:
            errors.append("no API key for fallback upload")
            return "", ""
        url = (self.base_url or "").rstrip("/") + "/api/v1/files/upload/" + flow_id
        try:
            r = await asyncio.to_thread(
                requests.post, url,
                headers={"x-api-key": key},
                files={"file": (name, data, XLSX_MIME)},
                timeout=120,
            )
            if r.status_code in (200, 201):
                path = ""
                try:
                    path = str((r.json() or {}).get("file_path") or "")
                except ValueError:
                    pass
                return "upload API", (path.split("/")[-1] if path else name)
            errors.append("upload API HTTP " + str(r.status_code) + ": " + r.text[:120])
        except Exception as e:
            errors.append("upload API: " + str(e)[:150])
        return "", ""

    async def _cleanup_store(self, svc, flow_id, errors):
        keep = _int(self.keep_days, 7)
        prefix = _safe(self.file_prefix) or "tracker"
        if keep <= 0:
            return
        try:
            names = await svc.list_files(flow_id=flow_id)
        except Exception as e:
            errors.append("cleanup list: " + str(e)[:100])
            return
        for n in names or []:
            n = str(n)
            if prefix in n and n.endswith(".xlsx") and _expired(n, keep):
                try:
                    await svc.delete_file(flow_id=flow_id, file_name=n)
                except Exception:
                    pass

    def _save_local(self, name, data, errors):
        folder = Path((self.local_dir or "").strip() or "/app/generated_reports")
        try:
            folder.mkdir(parents=True, exist_ok=True)
            path = folder / name
            path.write_bytes(data)
            keep = _int(self.keep_days, 7)
            prefix = _safe(self.file_prefix) or "tracker"
            for old in folder.glob("*" + prefix + "*.xlsx"):
                if _expired(old.name, keep):
                    try:
                        old.unlink()
                    except Exception:
                        pass
            return str(path)
        except Exception as e:
            errors.append("local save: " + str(e)[:150])
            return ""

    # ---------------- main ----------------
    async def _build(self):
        if getattr(self, "_res", None) is not None:
            return self._res
        errors = []
        payload = _data(self.tracker_data)
        if not payload:
            self._res = {"ok": False, "error": "No Tracker Data received. Connect Reviewer & Tracker > Tracker Data."}
            return self._res
        if payload.get("error"):
            self._res = {"ok": False, "error": "Upstream error: " + json.dumps(payload, default=str)[:400]}
            return self._res

        sheets, notes = _sheets_from(payload)
        tracker_rows = next((rows for n, c, rows in sheets if n == "Tracker"), [])
        agents = _agents(payload, tracker_rows)
        agent_part = _safe(agents[0], 40) if len(agents) == 1 else (str(len(agents)) + "_agents" if agents else "")
        dates = re.findall(r"\d{4}-\d{2}-\d{2}", str(payload.get("interval") or ""))
        date_part = (dates[0] + ("_to_" + dates[-1] if dates[-1] != dates[0] else "")) if dates else ""
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        name = "_".join(p for p in [_safe(self.file_prefix) or "tracker", agent_part,
                                    date_part.replace("-", ""), stamp] if p) + ".xlsx"

        try:
            data, counts = await asyncio.to_thread(
                _build_workbook, payload, sheets, notes, (self.interaction_url_template or "").strip(),
                bool(self.include_overview), bool(self.include_legend), max(0, _int(self.attention_rows, 25)))
        except ImportError:
            self._res = {"ok": False, "error": "openpyxl is not installed in MDS. Ask the MDS admins to add it."}
            return self._res
        except Exception as e:
            self._res = {"ok": False, "error": "Workbook build failed: " + str(e)[:300]}
            return self._res

        verdicts = _verdict_counts(payload, tracker_rows)
        info = {
            "ok": False,
            "file_name": name,
            "size_kb": round(len(data) / 1024, 1),
            "sheet_rows": counts,
            "agents": agents,
            "interval": payload.get("interval", ""),
            "answered_calls": payload.get("answered_calls", ""),
            "flagged_calls": payload.get("flagged_calls", len(tracker_rows)),
            "verdicts": verdicts,
            "qa_issues": payload.get("qa_issues", sum(1 for r in tracker_rows if _qa_issue(r))),
            "attention": [{k: r.get(k, "") for k in ATTENTION_COLS} for r in _attention(tracker_rows, 5)],
            "warnings": [str(w) for w in payload.get("warnings") or [] if w] + notes,
            "matcha_file": payload.get("file_name", "") if payload.get("uploaded_to_matcha") else "",
            "download_url": "",
            "local_path": "",
            "stored_via": "",
            "keep_days": _int(self.keep_days, 7),
            "errors": errors,
        }

        mode = str(self.storage or "")
        if mode.startswith("flow"):
            flow_id = self._flow_id()
            if not flow_id:
                errors.append("flow id not available - save the flow and run it again")
            else:
                via, stored = await self._save_store(flow_id, name, data, errors)
                if via:
                    info["stored_via"] = via
                    info["file_name"] = stored
                    info["download_url"] = ((self.base_url or "").rstrip("/") + "/api/v1/files/download/"
                                            + flow_id + "/" + urllib.parse.quote(stored))
                    info["ok"] = True
        if mode.endswith("local copy") or mode.startswith("local") or not info["ok"]:
            info["local_path"] = self._save_local(name, data, errors)
            if info["local_path"] and mode.startswith("local"):
                info["ok"] = True

        if self.inline_link and len(data) <= _int(self.inline_max_kb, 500) * 1024:
            info["inline_link"] = "data:" + XLSX_MIME + ";base64," + base64.b64encode(data).decode("ascii")

        self._res = info
        return info

    async def build_download(self) -> Message:
        r = await self._build()
        if not r.get("file_name"):
            self.status = "FAILED: " + str(r.get("error"))
            return Message(text="**Excel tracker not generated:** " + str(r.get("error")))

        lines = []
        if r.get("download_url"):
            lines += ["**Work Avoidance Tracker ready**: " + r["file_name"] + " (" + str(r["size_kb"]) + " KB)", "",
                      "[Download Excel tracker](" + r["download_url"] + ")", ""]
        elif r.get("local_path"):
            lines += ["**Tracker saved on the server only**: " + r["local_path"] + " (" + str(r["size_kb"]) + " KB)",
                      "No download link: the flow file store was not available.", ""]
        else:
            lines += ["**Tracker could not be saved.**", ""]

        lines.append("**Window:** " + str(r.get("interval") or "-"))
        lines.append("**Agents:** " + (", ".join(r.get("agents") or []) or "-"))
        lines.append("**Answered calls:** " + str(r.get("answered_calls", "-")) + " | **Flagged:** "
                     + str(r.get("flagged_calls", "-")) + " | **QA checks raised:** " + str(r.get("qa_issues", 0)))
        v = r.get("verdicts") or {}
        shown = [k for k in VERDICT_ORDER if v.get(k)] + [k for k in v if k not in VERDICT_ORDER and v.get(k)]
        if shown:
            lines += ["", "| Verdict | Calls |", "|---|---|"] + ["| " + k + " | " + str(v[k]) + " |" for k in shown]
        if r.get("attention"):
            lines += ["", "**Review first:**", "", "| Answered (ET) | Agent | Verdict | Risk | QA Check |",
                      "|---|---|---|---|---|"]
            for a in r["attention"]:
                lines.append("| " + " | ".join(str(a.get(k, "")).replace("|", "/")[:80] for k in
                                               ("Answered (ET)", "Agent", "AI Verdict", "AI Risk Score", "QA Check"))
                             + " |")
        if r.get("matcha_file"):
            lines += ["", "A copy was also uploaded to Matcha: " + r["matcha_file"]]
        if r.get("inline_link"):
            lines += ["", "[Download (inline copy)](" + r["inline_link"] + ")"]
        lines += ["", "The link requires an active MDS login. AI verdicts are advisory; a human reviewer decides."]
        if r.get("keep_days"):
            lines.append("Trackers are deleted automatically after " + str(r["keep_days"]) + " days.")
        if r.get("warnings"):
            lines += ["", "**Warnings:** " + " | ".join(r["warnings"])[:600]]
        if r.get("errors"):
            lines += ["", "Notes: " + "; ".join(r["errors"])[:400]]

        if r.get("ok"):
            self.status = "Saved " + r["file_name"] + " via " + (r.get("stored_via") or "local")
        else:
            self.status = "NOT SAVED: " + "; ".join(r.get("errors") or [])[:200]
        return Message(text="\n".join(lines))

    async def build_info(self) -> Data:
        r = await self._build()
        return Data(data={k: v for k, v in r.items() if k != "inline_link"})
